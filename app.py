import streamlit as st
import csv
import io
import openpyxl

def process_iccids(iccid_list):
    processed_data = []
    for iccid in iccid_list:
        iccid = iccid.strip()
        if not iccid:
            continue
        if len(iccid) > 19:
            processed = iccid[:19]
        else:
            processed = iccid
        processed_data.append((iccid, processed))
    return processed_data

def create_csv_string(data):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)
    for original, processed in data:
        writer.writerow([processed, processed])
    return output.getvalue()

def create_range_csv_string(start_iccids, end_iccids):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)
    max_length = max(len(start_iccids), len(end_iccids))
    for i in range(max_length):
        start = start_iccids[i] if i < len(start_iccids) else ""
        end = end_iccids[i] if i < len(end_iccids) else ""
        writer.writerow([start, end])
    return output.getvalue()

def read_iccids_from_upload(uploaded_file):
    filename = uploaded_file.name.lower()
    start_iccids = []
    end_iccids = []

    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        start_col = None
        end_col = None
        for idx, h in enumerate(headers):
            if 'start' in h and 'iccid' in h:
                start_col = idx
            elif 'end' in h and 'iccid' in h:
                end_col = idx

        if start_col is None or end_col is None:
            return None, None, "Could not find 'Start ICCID' and 'End ICCID' columns in the file."

        for row in ws.iter_rows(min_row=2, values_only=True):
            start_val = str(row[start_col]).strip() if row[start_col] is not None else ""
            end_val = str(row[end_col]).strip() if row[end_col] is not None else ""
            if start_val or end_val:
                start_iccids.append(start_val)
                end_iccids.append(end_val)

    elif filename.endswith('.csv'):
        content = uploaded_file.read().decode('utf-8')
        reader = csv.reader(io.StringIO(content), delimiter=',')
        headers_raw = next(reader, None)
        if not headers_raw:
            return None, None, "CSV file is empty."

        for delim in [',', ';', '\t']:
            test_reader = csv.reader(io.StringIO(content), delimiter=delim)
            test_headers = next(test_reader, None)
            if test_headers and len(test_headers) > 1:
                headers_raw = test_headers
                reader = csv.reader(io.StringIO(content), delimiter=delim)
                next(reader)
                break

        headers = [h.strip().lower() for h in headers_raw]
        start_col = None
        end_col = None
        for idx, h in enumerate(headers):
            if 'start' in h and 'iccid' in h:
                start_col = idx
            elif 'end' in h and 'iccid' in h:
                end_col = idx

        if start_col is None or end_col is None:
            return None, None, "Could not find 'Start ICCID' and 'End ICCID' columns in the CSV."

        for row in reader:
            if len(row) > max(start_col, end_col):
                start_val = row[start_col].strip()
                end_val = row[end_col].strip()
                if start_val or end_val:
                    start_iccids.append(start_val)
                    end_iccids.append(end_val)
    else:
        return None, None, "Unsupported file format. Please upload .xlsx or .csv files."

    return start_iccids, end_iccids, None

# Page configuration
st.set_page_config(
    page_title="ICCID Processor",
    page_icon="📱",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Custom CSS
st.markdown("""
    <style>
    .main {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 2rem;
    }
    .stApp > div {
        background: transparent;
    }
    .stTextArea textarea {
        border-radius: 12px !important;
        border: 2px solid #e0e0e0 !important;
        font-family: 'Courier New', monospace !important;
        font-size: 14px !important;
    }
    .stTextArea textarea:focus {
        border-color: #667eea !important;
        box-shadow: 0 0 0 2px rgba(102, 126, 234, 0.2) !important;
    }
    .stButton button {
        border-radius: 25px !important;
        font-weight: 600 !important;
        padding: 0.75rem 2rem !important;
        transition: all 0.3s ease !important;
        box-shadow: 0 4px 15px rgba(0, 0, 0, 0.2) !important;
    }
    .stButton button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(0, 0, 0, 0.3) !important;
    }
    .stDownloadButton button {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%) !important;
        color: white !important;
        border: none !important;
        border-radius: 25px !important;
        font-weight: 600 !important;
        padding: 0.75rem 2rem !important;
        box-shadow: 0 4px 15px rgba(102, 126, 234, 0.4) !important;
    }
    .stDownloadButton button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(102, 126, 234, 0.5) !important;
    }
    .stAlert {
        border-radius: 12px !important;
        animation: slideIn 0.3s ease-out;
    }
    @keyframes slideIn {
        from { opacity: 0; transform: translateY(-10px); }
        to { opacity: 1; transform: translateY(0); }
    }
    .streamlit-expanderHeader {
        border-radius: 8px !important;
        background-color: rgba(255, 255, 255, 0.1) !important;
    }
    .stDataFrame {
        border-radius: 12px !important;
        overflow: hidden !important;
    }
    .info-card {
        background: white;
        padding: 2rem;
        border-radius: 20px;
        box-shadow: 0 10px 40px rgba(0, 0, 0, 0.1);
        margin-bottom: 2rem;
    }
    .stat-card {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 1.5rem;
        border-radius: 15px;
        color: white;
        text-align: center;
        box-shadow: 0 8px 25px rgba(102, 126, 234, 0.3);
    }
    .stat-number {
        font-size: 2.5rem;
        font-weight: bold;
        margin: 0;
    }
    .stat-label {
        font-size: 1rem;
        opacity: 0.9;
        margin-top: 0.5rem;
    }
    h1 { color: white !important; text-shadow: 2px 2px 4px rgba(0, 0, 0, 0.2); }
    h2, h3 { color: white !important; }
    div[data-testid="column"] {
        background: white;
        padding: 2rem;
        border-radius: 20px;
        box-shadow: 0 10px 40px rgba(0, 0, 0, 0.1);
    }
    </style>
""", unsafe_allow_html=True)

# Header
st.markdown("<h1 style='text-align: center; font-size: 3rem; margin-bottom: 0.5rem;'>📱 ICCID Processor</h1>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; color: white; font-size: 1.2rem; margin-bottom: 3rem;'>Transform your ICCIDs with ease and precision</p>", unsafe_allow_html=True)

# Main content
col1, col2, col3 = st.columns([1, 8, 1])

with col2:
    tab1, tab2 = st.tabs(["🔧 Process ICCIDs", "📚 Instructions"])

    with tab1:
        st.markdown("<br>", unsafe_allow_html=True)

        mode = st.radio(
            "Choose processing mode:",
            [
                "📝 List Mode - Process multiple ICCIDs",
                "🔢 Range Mode - Generate from start to end",
                "📁 File Mode - Upload a CSV/Excel file",
            ],
            label_visibility="collapsed"
        )

        st.markdown("<br>", unsafe_allow_html=True)

        iccid_input = ""
        start_iccid = ""
        end_iccid = ""
        uploaded_file = None

        if mode == "📝 List Mode - Process multiple ICCIDs":
            st.markdown("### 📝 Enter Your ICCIDs")
            st.markdown("<p style='color: #666; margin-bottom: 1rem;'>Paste your ICCIDs below, one per line</p>", unsafe_allow_html=True)

            iccid_input = st.text_area(
                "ICCIDs",
                height=250,
                placeholder="8988228066623425355\n8988228066627262560\n8988228066627262660\n...",
                label_visibility="collapsed"
            )

            st.markdown("<br>", unsafe_allow_html=True)

            col_btn1, col_btn2, col_btn3 = st.columns([3, 2, 3])
            with col_btn2:
                generate_button = st.button("🚀 Generate CSV", type="primary", use_container_width=True)

        elif mode == "🔢 Range Mode - Generate from start to end":
            st.markdown("### 🔢 Enter ICCID Range")
            st.markdown("<p style='color: #666; margin-bottom: 1rem;'>Enter start and end ICCIDs (one per line, pairs will be matched)</p>", unsafe_allow_html=True)

            col_range1, col_range2 = st.columns(2)

            with col_range1:
                start_iccid = st.text_area(
                    "Start ICCIDs",
                    height=250,
                    placeholder="8988228066623425355\n8988228066623425365\n...",
                    help="Enter starting ICCIDs (one per line)"
                )

            with col_range2:
                end_iccid = st.text_area(
                    "End ICCIDs",
                    height=250,
                    placeholder="8988228066623425360\n8988228066623425370\n...",
                    help="Enter ending ICCIDs (one per line)"
                )

            st.markdown("<br>", unsafe_allow_html=True)

            col_btn1, col_btn2, col_btn3 = st.columns([3, 2, 3])
            with col_btn2:
                generate_button = st.button("🚀 Generate CSV", type="primary", use_container_width=True)

        else:
            st.markdown("### 📁 Upload ICCID File")
            st.markdown("<p style='color: #666; margin-bottom: 1rem;'>Upload a CSV or Excel file with <b>Start ICCID</b> and <b>End ICCID</b> columns</p>", unsafe_allow_html=True)

            uploaded_file = st.file_uploader(
                "Upload file",
                type=["xlsx", "xls", "csv"],
                label_visibility="collapsed",
                help="File must contain 'Start ICCID' and 'End ICCID' columns"
            )

            if uploaded_file:
                st.info(f"📄 Uploaded: **{uploaded_file.name}** ({uploaded_file.size / 1024:.1f} KB)")

            st.markdown("<br>", unsafe_allow_html=True)

            col_btn1, col_btn2, col_btn3 = st.columns([3, 2, 3])
            with col_btn2:
                generate_button = st.button("🚀 Generate CSV", type="primary", use_container_width=True)

        # Optional filename input
        st.markdown("<br>", unsafe_allow_html=True)
        custom_filename = st.text_input(
            "📁 Output filename (optional)",
            placeholder="iccids_output",
            help="Enter a custom name for the downloaded CSV file (without .csv extension)"
        )

        st.markdown("<br>", unsafe_allow_html=True)

        # Process and display results
        if generate_button:
            if mode == "📝 List Mode - Process multiple ICCIDs":
                default_filename = "iccids_output.csv"
                if not iccid_input.strip():
                    st.error("⚠️ Please enter at least one ICCID to process.")
                else:
                    iccids = [line.strip() for line in iccid_input.strip().split('\n') if line.strip()]

                    with st.spinner('Processing ICCIDs...'):
                        processed_data = process_iccids(iccids)

                    if not processed_data:
                        st.error("⚠️ No valid ICCIDs found. Please check your input.")
                    else:
                        csv_content = create_csv_string(processed_data)
                        output_filename = (custom_filename.strip() + ".csv") if custom_filename.strip() else default_filename

                        st.success(f"✅ Successfully processed {len(processed_data)} ICCID(s)!")

                        stat_col1, stat_col2, stat_col3 = st.columns(3)

                        with stat_col1:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(processed_data)}</p>
                                <p class='stat-label'>ICCIDs Processed</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col2:
                            trimmed_count = sum(1 for orig, proc in processed_data if len(orig) > 19)
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{trimmed_count}</p>
                                <p class='stat-label'>Trimmed</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col3:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(processed_data) - trimmed_count}</p>
                                <p class='stat-label'>Unchanged</p>
                            </div>
                            """, unsafe_allow_html=True)

                        st.markdown("<br>", unsafe_allow_html=True)

                        with st.expander("👀 Preview Processed Data", expanded=True):
                            preview_data = processed_data[:10]
                            preview_df = {
                                '📋 Original ICCID': [orig for orig, _ in preview_data],
                                '✨ Processed ICCID': [proc for _, proc in preview_data],
                                '📏 Length': [len(proc) for _, proc in preview_data]
                            }
                            st.dataframe(preview_df, use_container_width=True, height=400)
                            if len(processed_data) > 10:
                                st.info(f"📊 Showing 10 of {len(processed_data)} total rows")

                        st.markdown("<br>", unsafe_allow_html=True)

                        col_dl1, col_dl2, col_dl3 = st.columns([2, 4, 2])
                        with col_dl2:
                            st.download_button(
                                label="⬇️ Download CSV File",
                                data=csv_content,
                                file_name=output_filename,
                                mime="text/csv",
                                use_container_width=True
                            )

            elif mode == "🔢 Range Mode - Generate from start to end":
                default_filename = "iccids_range_output.csv"
                if not start_iccid.strip() or not end_iccid.strip():
                    st.error("⚠️ Please enter both start and end ICCIDs.")
                else:
                    start_iccids = [line.strip() for line in start_iccid.strip().split('\n') if line.strip()]
                    end_iccids = [line.strip() for line in end_iccid.strip().split('\n') if line.strip()]

                    if not start_iccids or not end_iccids:
                        st.error("⚠️ Please enter valid ICCIDs.")
                    else:
                        start_processed = [s[:19] if len(s) > 19 else s for s in start_iccids]
                        end_processed = [e[:19] if len(e) > 19 else e for e in end_iccids]
                        csv_content = create_range_csv_string(start_processed, end_processed)
                        output_filename = (custom_filename.strip() + ".csv") if custom_filename.strip() else default_filename

                        st.success("✅ ICCID range CSV generated successfully!")

                        stat_col1, stat_col2, stat_col3 = st.columns(3)

                        with stat_col1:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(start_processed)}</p>
                                <p class='stat-label'>Start ICCIDs</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col2:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(end_processed)}</p>
                                <p class='stat-label'>End ICCIDs</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col3:
                            total_pairs = max(len(start_processed), len(end_processed))
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{total_pairs}</p>
                                <p class='stat-label'>Total Rows</p>
                            </div>
                            """, unsafe_allow_html=True)

                        st.markdown("<br>", unsafe_allow_html=True)

                        with st.expander("👀 Preview Range Data", expanded=True):
                            preview_count = min(10, max(len(start_processed), len(end_processed)))
                            preview_starts = start_processed[:preview_count]
                            preview_ends = end_processed[:preview_count]
                            while len(preview_starts) < preview_count:
                                preview_starts.append("")
                            while len(preview_ends) < preview_count:
                                preview_ends.append("")
                            preview_df = {
                                '🎯 Start ICCID': preview_starts,
                                '🏁 End ICCID': preview_ends,
                            }
                            st.dataframe(preview_df, use_container_width=True, height=400)
                            total_rows = max(len(start_processed), len(end_processed))
                            if total_rows > 10:
                                st.info(f"📊 Showing 10 of {total_rows} total rows")
                            if len(start_processed) != len(end_processed):
                                st.warning(f"⚠️ Note: You have {len(start_processed)} start ICCIDs and {len(end_processed)} end ICCIDs. Empty cells will be added where needed.")

                        st.markdown("<br>", unsafe_allow_html=True)

                        col_dl1, col_dl2, col_dl3 = st.columns([2, 4, 2])
                        with col_dl2:
                            st.download_button(
                                label="⬇️ Download CSV File",
                                data=csv_content,
                                file_name=output_filename,
                                mime="text/csv",
                                use_container_width=True
                            )

            else:
                # File Mode
                default_filename = "iccids_file_output.csv"
                if uploaded_file is None:
                    st.error("⚠️ Please upload a file first.")
                else:
                    with st.spinner('Reading file...'):
                        start_iccids, end_iccids, error = read_iccids_from_upload(uploaded_file)

                    if error:
                        st.error(f"⚠️ {error}")
                    elif not start_iccids:
                        st.error("⚠️ No ICCID data found in the file.")
                    else:
                        start_processed = [s[:19] if len(s) > 19 else s for s in start_iccids]
                        end_processed = [e[:19] if len(e) > 19 else e for e in end_iccids]
                        csv_content = create_range_csv_string(start_processed, end_processed)
                        output_filename = (custom_filename.strip() + ".csv") if custom_filename.strip() else default_filename

                        st.success(f"✅ Successfully extracted {len(start_processed)} ICCID range(s) from file!")

                        stat_col1, stat_col2, stat_col3 = st.columns(3)

                        with stat_col1:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(start_processed)}</p>
                                <p class='stat-label'>Start ICCIDs</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col2:
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{len(end_processed)}</p>
                                <p class='stat-label'>End ICCIDs</p>
                            </div>
                            """, unsafe_allow_html=True)

                        with stat_col3:
                            total_pairs = max(len(start_processed), len(end_processed))
                            st.markdown(f"""
                            <div class='stat-card'>
                                <p class='stat-number'>{total_pairs}</p>
                                <p class='stat-label'>Total Rows</p>
                            </div>
                            """, unsafe_allow_html=True)

                        st.markdown("<br>", unsafe_allow_html=True)

                        with st.expander("👀 Preview Extracted Data", expanded=True):
                            preview_count = min(10, max(len(start_processed), len(end_processed)))
                            preview_starts = start_processed[:preview_count]
                            preview_ends = end_processed[:preview_count]
                            while len(preview_starts) < preview_count:
                                preview_starts.append("")
                            while len(preview_ends) < preview_count:
                                preview_ends.append("")
                            preview_df = {
                                '🎯 Start ICCID': preview_starts,
                                '🏁 End ICCID': preview_ends,
                            }
                            st.dataframe(preview_df, use_container_width=True, height=400)
                            total_rows = max(len(start_processed), len(end_processed))
                            if total_rows > 10:
                                st.info(f"📊 Showing 10 of {total_rows} total rows")

                        st.markdown("<br>", unsafe_allow_html=True)

                        col_dl1, col_dl2, col_dl3 = st.columns([2, 4, 2])
                        with col_dl2:
                            st.download_button(
                                label="⬇️ Download CSV File",
                                data=csv_content,
                                file_name=output_filename,
                                mime="text/csv",
                                use_container_width=True
                            )

    with tab2:
        st.markdown("<br>", unsafe_allow_html=True)

        st.markdown("### 🎯 How to Use")

        st.markdown("""
        #### 📝 List Mode
        Process multiple ICCIDs at once:
        - Paste your ICCIDs (one per line)
        - Click "Generate CSV"
        - Get a CSV with duplicate columns for each processed ICCID

        #### 🔢 Range Mode
        Generate a CSV with start and end ICCIDs:
        - Enter starting ICCIDs in the left column
        - Enter ending ICCIDs in the right column
        - Click "Generate CSV"
        - Get a CSV with start in first column, end in second column

        #### 📁 File Mode
        Upload a CSV or Excel file:
        - File must have **Start ICCID** and **End ICCID** columns
        - Other columns (Order no., Itemcode, etc.) are ignored
        - ICCIDs are extracted and trimmed automatically

        ---
        """)

        col_inst1, col_inst2 = st.columns(2)

        with col_inst1:
            st.markdown("""
            #### Step 1: Input
            - Choose your processing mode
            - Enter ICCIDs or upload a file
            - Empty lines will be ignored

            #### Step 2: Process
            - Optionally set a custom output filename
            - Click the "Generate CSV" button
            - See statistics and preview
            """)

        with col_inst2:
            st.markdown("""
            #### Processing Rules
            - **ICCIDs > 19 digits**: Trimmed to 19 digits
            - **ICCIDs ≤ 19 digits**: Kept unchanged
            - **Output format**: CSV with semicolon delimiter

            #### Step 3: Download
            - Review the preview
            - Click "Download CSV File"
            - Save to your device
            """)

        st.markdown("<br>", unsafe_allow_html=True)

        st.markdown("### 💡 Example Input")
        st.code("""8988228066623425355
8988228066627262560
8988228066627262660
8988228066627262760
8988228066627262860""", language=None)

        st.markdown("### 📤 Example Output")
        st.code("""8988228066623425355;8988228066623425355
8988228066627262560;8988228066627262560
8988228066627262660;8988228066627262660""", language="csv")

        st.markdown("<br>", unsafe_allow_html=True)

        st.markdown("### 📁 File Mode Template")
        st.markdown("Your uploaded file should have columns like this:")
        st.code("""Order no. | Itemcode | Description | Qty. | Start ICCID         | End ICCID
NEXD-2052 | 220003   | IoT SIM...  | 500  | 8988228066632854560 | 8988228066632855059
NEXD-2052 | 220003   | IoT SIM...  | 500  | 8988228066632855060 | 8988228066632855559""", language=None)

# Footer
st.markdown("<br><br>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; color: white; opacity: 0.7;'>Made by Othmane with ❤️ for efficient ICCID processing</p>", unsafe_allow_html=True)
